"""Generate binary/encoding fixtures for spec 007 (T002). Run once: uv run python bench/fixtures/extraction/_generate.py"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).resolve().parent

# --- windows-1251 encoded page -------------------------------------------------
cp1251 = (
    "<!DOCTYPE html><html lang=\"ru\"><head><meta charset=\"windows-1251\">"
    "<title>Иванов Сергей Петрович</title></head><body><article><h1>Иванов Сергей Петрович</h1>"
    "<p>Родился в Казани, живёт в Москве. Контакт: sergey@example-ivanov.ru</p></article>"
    "<footer>Навигация: Главная · О сайте</footer></body></html>"
)
(HERE / "cp1251_page.html").write_bytes(cp1251.encode("windows-1251"))

# --- corrupt binary ------------------------------------------------------------
(HERE / "corrupt.bin").write_bytes(b"\x00\x01\x02PDF-like garbage \xff\xfe" + b"\x9c" * 128)

# --- PDF ------------------------------------------------------------------------
try:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_metadata(
        {"/Author": "Sergey Ivanov", "/Producer": "cognitive-fixtures", "/Title": "Ivanov CV"}
    )
    with (HERE / "sample.pdf").open("wb") as fh:
        writer.write(fh)
except Exception as exc:  # pragma: no cover - generation-time only
    print(f"pdf skipped: {exc}")

# --- DOCX -----------------------------------------------------------------------
try:
    import docx

    d = docx.Document()
    d.add_paragraph("Sergey Ivanov — research engineer, born in Kazan, lives in Moscow.")
    d.add_paragraph("Contact: sergey@example-ivanov.ru")
    d.core_properties.author = "Sergey Ivanov"
    d.save(HERE / "sample.docx")
except Exception as exc:  # pragma: no cover
    print(f"docx skipped: {exc}")

# --- JPEG with EXIF GPS (Kazan approx: 55.7963 N, 49.1088 E) --------------------
try:
    import struct

    from PIL import Image

    def _dms(lat: bool) -> tuple[tuple[int, int], tuple[int, int], tuple[int, int]]:
        value = 55.7963 if lat else 49.1088
        deg = int(abs(value))
        minutes = int((abs(value) - deg) * 60)
        sec = ((abs(value) - deg) * 60 - minutes) * 60
        return (deg, 1), (minutes, 1), (int(sec * 1_000_000), 1_000_000)

    LAT_REF, LON_REF = b"N\0", b"E\0"

    TIFF_HEADER = b"II\x2a\x00" + struct.pack("<I", 8)

    def _rationals_entry(tag: int, values: tuple[tuple[int, int], ...], ifd_start: int, value_base: int) -> tuple[bytes, bytes]:
        """EXIF RATIONAL entry: count reflects the number of rationals (dms → 3)."""
        return (
            struct.pack("<HHI", tag, 0x0005, len(values)) + struct.pack("<I", value_base),
            b"".join(struct.pack("<II", n, d) for n, d in values),
        )

    def _exif_blob() -> bytes:
        gps: list[tuple[tuple[int, int], tuple[int, int], tuple[int, int]]] = [_dms(True), _dms(False)]
        make, model = b"CogCam\0", b"CogCam One\0"
        ifd0_size = 2 + 3 * 12 + 4
        # GPS IFD is placed right after IFD0 + both ASCII strings (offsets in TIFF coords).
        gps_ifd_start = 8 + ifd0_size + len(make) + len(model)
        gps_count = 4
        gps_ifd_size = 2 + gps_count * 12 + 4
        rational_base = gps_ifd_start + gps_ifd_size
        gps_entries = bytearray()
        gps_entries += struct.pack("<H", gps_count)
        gps_entries += struct.pack("<HHI", 1, 2, 2) + LAT_REF.ljust(4, b"\x00")
        gps_entries += _rationals_entry(2, gps[0], 0, rational_base)[0]
        gps_entries += struct.pack("<HHI", 3, 2, 2) + LON_REF.ljust(4, b"\x00")
        gps_entries += _rationals_entry(4, gps[1], 0, rational_base + 24)[0]
        gps_entries += struct.pack("<I", 0)
        gps_body = b"".join(struct.pack("<II", n, d) for pair in gps for n, d in pair)

        zero_base = 8 + ifd0_size
        zero_entry_make = struct.pack("<HHI", 0x010F, 2, len(make)) + struct.pack("<I", zero_base)
        zero_entry_model = struct.pack("<HHI", 0x0110, 2, len(model)) + struct.pack("<I", zero_base + len(make))
        zero_entry_gps = struct.pack("<HHI", 0x8825, 0x0004, 1) + struct.pack("<I", gps_ifd_start)
        zero_ifd = (
            struct.pack("<H", 3)
            + zero_entry_make
            + zero_entry_model
            + zero_entry_gps
            + struct.pack("<I", 0)
        )
        return b"Exif\0\0" + TIFF_HEADER + zero_ifd + make + model + bytes(gps_entries) + gps_body

    img = Image.new("RGB", (64, 64), color=(30, 60, 90))
    img.save(HERE / "sample_exif.jpg", exif=_exif_blob(), quality=85)
except Exception as exc:  # pragma: no cover
    print(f"jpeg skipped: {exc}")

print("fixtures done")
